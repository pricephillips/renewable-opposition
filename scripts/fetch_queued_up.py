"""Store LBNL Queued Up: every solar, wind and storage project that sought grid interconnection.

Queued Up (Lawrence Berkeley National Laboratory and GridTracker,
https://emp.lbl.gov/queues; CC BY 4.0) compiles the interconnection queues of
the seven ISOs and RTOs and about 50 other utilities, together about 98% of
installed U.S. capacity: one row per interconnection request, active,
withdrawn, suspended or operational, with its request, agreement, withdrawal
and commercial operation dates. It is the universe of proposed projects this
repo lacks (docs/predictive_foundation_design.md, section 5): projects are in
it whether or not anyone objected to them.

The workbook's "Complete Queue Data" sheet is read, and the requests with a
solar, wind (onshore or offshore) or storage component are kept, one row each,
in data/reference/queue_projects.csv. Hybrids are kept whole: technology lists
the solar, wind and storage parts, and other_mw holds any other part (the gas
in "Gas+Battery").

Placement. The workbook gives one county per request (for a request spanning
several, only the first; see its codebook) and that county's FIPS code. A code
that is a 2024 county in the request's state is used as given
(county_fips_method source_fips). Otherwise the county name is looked up as
the build looks up a contested project's county (classify.county_fips_all,
method name). Anything else stays unplaced, with the reason in county_miss:
no_state, no_county, not_a_2024_county (Connecticut's retired counties, which
the queues still report) or not_in_lookup. Nothing is guessed.

The workbook itself is not committed (about 15 MB). data/reference/
queued_up_manifest.csv records its URL, edition, SHA-256 and the rows read,
kept and dropped by reason. The link is found on LBNL's pages (the queues page
links each edition's publication page, which links the workbook); those pages
refuse some automated clients, and when they do the last workbook URL in the
manifest is fetched instead. A workbook whose SHA-256 matches the manifest is
still read through the schema guard, but nothing is rewritten, so an
unchanged edition commits nothing.

A workbook whose sheet or columns differ from what this script reads is
refused and nothing is written (the schema guard).

Usage
  python scripts/fetch_queued_up.py                   find, fetch, write
  python scripts/fetch_queued_up.py --file FILE.xlsx  read a local copy
  python scripts/fetch_queued_up.py --dry-run         print the plan
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

import requests

sys.path.insert(0, str(Path(__file__).resolve().parent))
import classify  # noqa: E402
import geo  # noqa: E402
from common import ROOT, STATE_CODES, STATE_NAMES, read_csv, write_csv  # noqa: E402

REFERENCE_DIR = ROOT / "data" / "reference"
QUEUE_PATH = REFERENCE_DIR / "queue_projects.csv"
MANIFEST_PATH = REFERENCE_DIR / "queued_up_manifest.csv"
FIPS_LOOKUP = ROOT / "data" / "county_fips_lookup.json"

QUEUES_PAGE = "https://emp.lbl.gov/queues"
# The 2026 edition (data through 2025), used when the pages refuse this
# client and the manifest has no URL yet.
KNOWN_URL = "https://eta-publications.lbl.gov/sites/default/files/2026-05/lbnl_ix_queue_data_file_thru2025.xlsx"
WORKBOOK_LINK = re.compile(r"https://eta-publications\.lbl\.gov/sites/default/files/\d{4}-\d{2}/"
                           r"[\w.-]*queue[\w.-]*thru(\d{4})\.xlsx", re.I)
EDITION_LINK = re.compile(r'href="((?:https://emp\.lbl\.gov)?/publications/queued-(\d{4})-edition[\w-]*)"', re.I)
USER_AGENT = ("renewable-opposition-pipeline/0.1 "
              "(research; contact: github.com/pricephillips/renewable-opposition)")
LICENSE = "CC BY 4.0 (Lawrence Berkeley National Laboratory and GridTracker)"

SHEET_SUFFIX = "Complete Queue Data"
REQUIRED = ["q_id", "q_status", "q_date", "prop_date", "on_date", "wd_date", "ia_date", "IA_phase_clean",
            "county", "state", "fips_code", "poi_name", "region", "project_name", "utility", "entity",
            "developer", "cluster", "service", "project_type", "type_1", "type_2", "type_3", "type_clean",
            "mw_1", "mw_2", "mw_3"]
# Workbook resource type -> the shared technology vocabulary. A request is
# kept when any of its types is here.
KEEP_TYPES = {"Solar": "solar", "Wind": "wind", "Offshore Wind": "wind", "Battery": "battery_storage",
              "Other Storage": "battery_storage"}
MW_COLUMN = {"solar": "solar_mw", "wind": "wind_mw", "battery_storage": "storage_mw"}

FIELDS = ["queue_project_id", "edition", "region", "entity", "utility", "q_id", "project_name", "developer",
          "status", "ia_phase", "project_type", "service", "type_clean", "technology", "offshore",
          "capacity_mw", "solar_mw", "wind_mw", "storage_mw", "other_mw", "request_date",
          "proposed_online_date", "ia_date", "withdrawn_date", "operational_date", "state", "county",
          "source_fips", "county_fips", "county_fips_method", "county_miss", "poi_name", "cluster"]
MANIFEST_FIELDS = ["source", "release", "url", "file_name", "sha256", "bytes", "fetched_at", "found_by",
                   "rows_read", "rows_kept", "rows_dropped", "license"]


class SchemaError(ValueError):
    pass


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _s(v) -> str:
    """A cell as one line of text: the workbook has line breaks inside some
    names, which would split a row for line-based tools."""
    return "" if v is None else " ".join(str(v).split())


def _num(v) -> float | None:
    try:
        f = float(_s(v))
    except ValueError:
        return None
    return f


def _mw(v: float) -> str:
    return f"{v:.3f}".rstrip("0").rstrip(".") if v else ""


def iso(v) -> str:
    """A workbook cell as YYYY-MM-DD, or "" when it holds no date."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", _s(v))
    return m.group(0) if m else ""


def slug(v: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", v.lower()).strip("-")


def edition_of(url: str) -> str:
    m = re.search(r"thru(\d{4})", url)
    return f"Queued Up, data through {m.group(1)}" if m else "Queued Up"


# ── Finding the workbook ────────────────────────────────────────────────────

def workbook_links(html: str) -> list[tuple[int, str]]:
    """(data year, URL) of every workbook linked on a page, newest first."""
    return sorted({(int(m.group(1)), m.group(0)) for m in WORKBOOK_LINK.finditer(html)}, reverse=True)


def edition_pages(html: str) -> list[str]:
    """Publication pages for each edition linked from the queues page, newest first."""
    found = {(int(m.group(2)), requests.compat.urljoin(QUEUES_PAGE, m.group(1))) for m in EDITION_LINK.finditer(html)}
    return [u for _, u in sorted(found, reverse=True)]


def last_known_url() -> str:
    rows = read_csv(MANIFEST_PATH) if MANIFEST_PATH.exists() else []
    return rows[-1]["url"] if rows and rows[-1].get("url") else KNOWN_URL


def find_workbook(session: requests.Session) -> tuple[str, str]:
    """(URL, how it was found). The queues page, then the newest edition's
    publication page; the last known URL when the pages refuse this client."""
    try:
        r = session.get(QUEUES_PAGE, timeout=60)
        if r.ok:
            links = workbook_links(r.text)
            if links:
                return links[0][1], "queues page"
            for page in edition_pages(r.text)[:2]:
                p = session.get(page, timeout=60)
                if p.ok and workbook_links(p.text):
                    return workbook_links(p.text)[0][1], "edition page"
    except requests.RequestException:
        pass
    return last_known_url(), "last known URL (LBNL pages not readable by this client)"


# ── Reading the workbook ────────────────────────────────────────────────────

def read_workbook(data: bytes) -> tuple[list[dict], int]:
    """(request rows as dicts keyed by the codebook's field names, rows read)."""
    import openpyxl
    if not data.startswith(b"PK"):
        raise SchemaError("not an xlsx workbook")
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = [s for s in wb.sheetnames if s.strip().endswith(SHEET_SUFFIX)]
    if len(sheets) != 1:
        raise SchemaError(f"expected one sheet ending {SHEET_SUFFIX!r}, found {sheets}")
    rows = wb[sheets[0]].iter_rows(values_only=True)
    header = None
    for _ in range(10):
        r = next(rows, None)
        if r and "q_id" in r:
            header = [_s(c) for c in r]
            break
    if header is None:
        raise SchemaError(f"{sheets[0]}: no header row with q_id in the first ten rows")
    lacks = [c for c in REQUIRED if c not in header]
    if lacks:
        raise SchemaError(f"{sheets[0]}: missing column(s) {lacks}")
    out = [dict(zip(header, r)) for r in rows if r and any(c not in (None, "") for c in r)]
    if not out:
        raise SchemaError(f"{sheets[0]}: no rows")
    return out, len(out)


def place(state: str, county: str, fips_code, lookup: dict[str, str]) -> tuple[str, str, str, str]:
    """(source_fips, county_fips, county_fips_method, county_miss)."""
    src = ""
    n = _num(fips_code)
    if n is not None and n > 0:
        src = str(int(n)).zfill(5)
    if state not in STATE_CODES:
        return src, "", "", "no_state"
    st_fips = next((f for f, code in STATE_FIPS.items() if code == state), "")
    if src and src[:2] == st_fips and geo.known(src) and src not in classify.NOT_2024_COUNTY:
        return src, src, "source_fips", ""
    if not county:
        return src, "", "", "no_county"
    row = {"state": state, "county": county.replace("/", ", ")}
    found, method = classify.county_fips_all("contested_projects", row, lookup, STATE_NAMES)
    if len(found) == 1:
        return src, found[0], method, ""
    _, reason = classify.county_fips("contested_projects", row, lookup, STATE_NAMES)
    return src, "", "", reason or "not_in_lookup"


STATE_FIPS = {
    "01": "AL", "02": "AK", "04": "AZ", "05": "AR", "06": "CA", "08": "CO", "09": "CT", "10": "DE",
    "11": "DC", "12": "FL", "13": "GA", "15": "HI", "16": "ID", "17": "IL", "18": "IN", "19": "IA",
    "20": "KS", "21": "KY", "22": "LA", "23": "ME", "24": "MD", "25": "MA", "26": "MI", "27": "MN",
    "28": "MS", "29": "MO", "30": "MT", "31": "NE", "32": "NV", "33": "NH", "34": "NJ", "35": "NM",
    "36": "NY", "37": "NC", "38": "ND", "39": "OH", "40": "OK", "41": "OR", "42": "PA", "44": "RI",
    "45": "SC", "46": "SD", "47": "TN", "48": "TX", "49": "UT", "50": "VT", "51": "VA", "53": "WA",
    "54": "WV", "55": "WI", "56": "WY", "72": "PR",
}


def project_rows(raw: list[dict], edition: str, lookup: dict[str, str]) -> tuple[list[dict], Counter]:
    """(kept rows, dropped rows by reason)."""
    out, dropped = [], Counter()
    for r in raw:
        types = [_s(r.get(f"type_{k}")) for k in (1, 2, 3)]
        mws = [_num(r.get(f"mw_{k}")) or 0.0 for k in (1, 2, 3)]
        tech = sorted({KEEP_TYPES[t] for t in types if t in KEEP_TYPES})
        if not tech:
            dropped["no_solar_wind_or_storage"] += 1
            continue
        state = _s(r.get("state")).upper()
        if state and state not in STATE_CODES:
            dropped["outside_the_us_states"] += 1
            continue
        by_tech = Counter()
        other = 0.0
        for t, mw in zip(types, mws):
            if t in KEEP_TYPES:
                by_tech[MW_COLUMN[KEEP_TYPES[t]]] += mw
            elif t:
                other += mw
        county = _s(r.get("county"))
        src, fips, method, miss = place(state, county, r.get("fips_code"), lookup)
        q_id = _s(r.get("q_id"))
        out.append({
            "edition": edition, "region": _s(r.get("region")), "entity": _s(r.get("entity")),
            "utility": _s(r.get("utility")), "q_id": q_id, "project_name": _s(r.get("project_name")),
            "developer": _s(r.get("developer")), "status": _s(r.get("q_status")).lower(),
            "ia_phase": _s(r.get("IA_phase_clean")), "project_type": _s(r.get("project_type")),
            "service": _s(r.get("service")), "type_clean": _s(r.get("type_clean")),
            "technology": ";".join(tech), "offshore": "yes" if "Offshore Wind" in types else "",
            "capacity_mw": _mw(sum(mws)), "solar_mw": _mw(by_tech["solar_mw"]),
            "wind_mw": _mw(by_tech["wind_mw"]), "storage_mw": _mw(by_tech["storage_mw"]),
            "other_mw": _mw(other), "request_date": iso(r.get("q_date")),
            "proposed_online_date": iso(r.get("prop_date")), "ia_date": iso(r.get("ia_date")),
            "withdrawn_date": iso(r.get("wd_date")), "operational_date": iso(r.get("on_date")),
            "state": state, "county": county, "source_fips": src, "county_fips": fips,
            "county_fips_method": method, "county_miss": miss, "poi_name": _s(r.get("poi_name")),
            "cluster": _s(r.get("cluster")),
        })
    assign_ids(out)
    out.sort(key=lambda p: (p["state"], p["entity"], p["queue_project_id"]))
    return out, dropped


def assign_ids(rows: list[dict]) -> None:
    """<entity>:<q_id>, unique within the file. A queue id that is blank or
    repeated within its entity gets a suffix from the request's own fields."""
    base = [f"{slug(p['entity']) or 'unknown'}:{slug(p['q_id']) or 'none'}" for p in rows]
    counts = Counter(base)
    for p, b in zip(rows, base):
        if counts[b] == 1 and not b.endswith(":none") and not b.endswith(":not-assigned"):
            p["queue_project_id"] = b
            continue
        key = "|".join(p[k] for k in ("request_date", "state", "county", "type_clean", "capacity_mw",
                                      "poi_name", "project_name", "status"))
        p["queue_project_id"] = f"{b}~{hashlib.sha1(key.encode()).hexdigest()[:8]}"
    seen = Counter(p["queue_project_id"] for p in rows)
    for i, p in enumerate(r for r in rows if seen[r["queue_project_id"]] > 1):
        p["queue_project_id"] += f"~{i + 1}"


def build(url: str, data: bytes, found_by: str, lookup: dict[str, str]) -> tuple[list[dict], dict]:
    raw, read = read_workbook(data)
    edition = edition_of(url)
    rows, dropped = project_rows(raw, edition, lookup)
    manifest = {"source": "queued_up", "release": edition, "url": url, "file_name": url.rsplit("/", 1)[-1],
                "sha256": sha256(data), "bytes": len(data),
                "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "found_by": found_by,
                "rows_read": read, "rows_kept": len(rows),
                "rows_dropped": "; ".join(f"{k} {v}" for k, v in sorted(dropped.items())), "license": LICENSE}
    return rows, manifest


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--file", type=Path, help="a local copy of the workbook instead of downloading")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    if args.dry_run:
        print(f"[dry-run] find the workbook on {QUEUES_PAGE}, else {last_known_url()}")
        print(f"[dry-run] would write {QUEUE_PATH.relative_to(ROOT)} and {MANIFEST_PATH.relative_to(ROOT)}")
        return 0
    if args.file:
        url, data, found_by = last_known_url(), args.file.read_bytes(), f"local file {args.file.name}"
    else:
        session = requests.Session()
        session.headers["User-Agent"] = USER_AGENT
        url, found_by = find_workbook(session)
        try:
            r = session.get(url, timeout=300)
            r.raise_for_status()
        except requests.RequestException as exc:
            print(f"Download failed, nothing written: {exc}", file=sys.stderr)
            return 1
        data = r.content
    lookup = json.loads(FIPS_LOOKUP.read_text(encoding="utf-8"))
    try:
        rows, manifest = build(url, data, found_by, lookup)
    except SchemaError as exc:
        print(f"Refused by the schema guard, nothing written: {exc}", file=sys.stderr)
        return 1
    print(f"{manifest['release']}: {url} ({found_by}), {manifest['rows_read']} requests read, "
          f"{manifest['rows_kept']} kept; dropped: {manifest['rows_dropped'] or 'none'}")
    status = Counter(r["status"] for r in rows)
    print("Status: " + ", ".join(f"{k} {v}" for k, v in status.most_common()))
    method = Counter(r["county_fips_method"] or f"unplaced ({r['county_miss']})" for r in rows)
    print("Placement: " + ", ".join(f"{k} {v}" for k, v in method.most_common()))
    previous = read_csv(MANIFEST_PATH) if MANIFEST_PATH.exists() else []
    if previous and previous[-1].get("sha256") == manifest["sha256"] and QUEUE_PATH.exists():
        print(f"Unchanged workbook (SHA-256 {manifest['sha256'][:12]}); nothing rewritten")
        return 0
    write_csv(QUEUE_PATH, rows, FIELDS)
    write_csv(MANIFEST_PATH, [manifest], MANIFEST_FIELDS)
    print(f"Wrote {QUEUE_PATH.relative_to(ROOT)} ({len(rows)} rows)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
