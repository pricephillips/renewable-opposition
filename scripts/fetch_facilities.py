"""Store the federal inventories of built, planned and cancelled power plants.

Three public-domain datasets record what happened to a renewable project after
the fight over it: whether a plant was built, is planned or under construction,
or was cancelled.

  eia860m  EIA-860M, Preliminary Monthly Electric Generator Inventory
           (https://www.eia.gov/electricity/data/eia860m/). The newest monthly
           workbook: its Operating, Planned, Retired and "Canceled or
           Postponed" tabs, every generator of 1 MW or more.
  uswtdb   U.S. Wind Turbine Database (USGS, LBNL and ACP;
           https://energy.usgs.gov/uswtdb/), every built land-based and
           offshore turbine, located from imagery.
  uspvdb   U.S. Large-Scale Solar Photovoltaic Database (USGS and LBNL;
           https://energy.usgs.gov/uspvdb/), every ground-mounted PV
           facility of 1 MW or more, digitized from imagery.

All three are U.S. government works in the public domain. Only solar, wind,
battery storage and geothermal plants are kept, one row per plant and status
(EIA), project (USWTDB, turbines grouped by project name, state and EIA plant
id) or facility (USPVDB), in data/reference/facilities.csv. Each plant is
placed in its county by its coordinates (geo.county_at); a plant with no
usable coordinates keeps the county its source names, unplaced.

The source files themselves are not committed (the EIA workbook is about 14 MB
a month). data/reference/facilities_manifest.csv records, for each, the URL,
the release, the SHA-256 and size of the file read and the rows read and kept,
so the extract can be traced to the exact file. EIA keeps every monthly
workbook in its archive under the same name.

A file whose sheets or headers differ from what this script reads is refused
and nothing is written (the schema guard, as for the NREL spreadsheets).

Usage
  python scripts/fetch_facilities.py                       fetch all three
  python scripts/fetch_facilities.py --eia FILE --uswtdb FILE --uspvdb FILE
                                                           read local copies
  python scripts/fetch_facilities.py --dry-run             print the plan
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import requests

sys.path.insert(0, str(Path(__file__).resolve().parent))
import geo  # noqa: E402
from common import ROOT, STATE_CODES, write_csv  # noqa: E402

REFERENCE_DIR = ROOT / "data" / "reference"
FACILITIES_PATH = REFERENCE_DIR / "facilities.csv"
MANIFEST_PATH = REFERENCE_DIR / "facilities_manifest.csv"

EIA_INDEX = "https://www.eia.gov/electricity/data/eia860m/"
USWTDB_URL = "https://energy.usgs.gov/uswtdb/assets/data/uswtdbCSV.zip"
USPVDB_URL = "https://energy.usgs.gov/uspvdb/assets/data/uspvdbCSV.zip"
USER_AGENT = ("renewable-opposition-pipeline/0.1 "
              "(research; contact: github.com/pricephillips/renewable-opposition)")

FIELDS = ["facility_id", "source", "release", "plant_id", "eia_plant_id", "name", "operator",
          "state", "county", "county_fips", "technology", "status", "status_detail",
          "capacity_mw", "year", "month", "units", "latitude", "longitude"]
MANIFEST_FIELDS = ["source", "release", "url", "file_name", "sha256", "bytes", "fetched_at",
                   "rows_read", "rows_kept"]

# EIA technology -> the shared vocabulary. Anything else is not kept.
EIA_TECH = {
    "Solar Photovoltaic": "solar",
    "Solar Thermal with Energy Storage": "solar",
    "Solar Thermal without Energy Storage": "solar",
    "Onshore Wind Turbine": "wind",
    "Offshore Wind Turbine": "wind",
    "Batteries": "battery_storage",
    "Geothermal": "geothermal",
}
# Planned-tab status codes that mean construction has started.
UNDER_CONSTRUCTION = {"U", "V", "TS"}
# Sheet -> status, for the sheets read. The Puerto Rico tabs have the same
# layout and are read the same way.
EIA_SHEETS = {
    "Operating": "operating", "Planned": "planned", "Retired": "retired",
    "Canceled or Postponed": "canceled_or_postponed",
    "Operating_PR": "operating", "Planned_PR": "planned", "Retired_PR": "retired",
}
EIA_REQUIRED = ["Entity Name", "Plant ID", "Plant Name", "Plant State", "County", "Generator ID",
                "Nameplate Capacity (MW)", "Technology", "Latitude", "Longitude"]
EIA_DATE_COLS = {"operating": ("Operating Year", "Operating Month"),
                 "planned": ("Planned Operation Year", "Planned Operation Month"),
                 "retired": ("Retirement Year", "Retirement Month")}
USWTDB_REQUIRED = ["case_id", "eia_id", "t_state", "t_county", "t_fips", "p_name", "p_year",
                   "p_tnum", "p_cap", "t_offshore", "xlong", "ylat"]
USPVDB_REQUIRED = ["case_id", "eia_id", "p_state", "p_county", "ylat", "xlong", "p_name",
                   "p_year", "p_cap_ac", "p_battery"]


class SchemaError(ValueError):
    pass


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _num(v) -> float | None:
    try:
        f = float(str(v).strip())
    except (TypeError, ValueError):
        return None
    return f


def _int_text(v) -> str:
    f = _num(v)
    return str(int(f)) if f is not None else ""


def _mw(total: float) -> str:
    return f"{total:.1f}".rstrip("0").rstrip(".")


def _place(lat: float | None, lon: float | None) -> str:
    if lat is None or lon is None or not (-180 <= lon <= -60 and 15 <= lat <= 72):
        return ""
    return geo.county_at(lat, lon)


# ── EIA-860M ────────────────────────────────────────────────────────────────

def eia_candidates(index_html: str) -> list[str]:
    """Workbook URLs on the EIA-860M page, newest first. The page lists months
    whose files are not out yet (they return an HTML page), so the caller
    tries each in turn."""
    months = ["january", "february", "march", "april", "may", "june", "july", "august",
              "september", "october", "november", "december"]
    found = {}
    for href in re.findall(r'href="([^"]*?([a-z]+)_generator(\d{4})\.xlsx)"', index_html):
        url, month, year = href
        if month in months:
            found[(int(year), months.index(month))] = requests.compat.urljoin(EIA_INDEX, url)
    return [found[k] for k in sorted(found, reverse=True)]


def read_eia(data: bytes) -> tuple[str, list[dict], int]:
    """(release, plant rows, generator rows read) from an EIA-860M workbook."""
    import openpyxl
    if not data.startswith(b"PK"):
        raise SchemaError("not an xlsx workbook")
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    missing = [s for s in ("Operating", "Planned", "Retired", "Canceled or Postponed")
               if s not in wb.sheetnames]
    if missing:
        raise SchemaError(f"EIA-860M workbook lacks sheet(s) {missing}")
    release = ""
    plants: dict[tuple, dict] = {}
    read = 0
    for sheet, status in EIA_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        title = str(next(rows)[0] or "")
        m = re.search(r"as of (\w+ \d{4})", title)
        if m and not release:
            release = f"EIA-860M {m.group(1)}"
        header = None
        for r in rows:
            if r and "Entity ID" in r:
                header = [str(c or "").strip() for c in r]
                break
        if header is None:
            raise SchemaError(f"{sheet}: no header row")
        lacks = [c for c in EIA_REQUIRED if c not in header]
        if status == "planned":
            lacks += [c for c in ("Status",) if c not in header]
        for c in EIA_DATE_COLS.get(status, ()):
            if c not in header:
                lacks.append(c)
        if lacks:
            raise SchemaError(f"{sheet}: missing column(s) {lacks}")
        ix = {c: i for i, c in enumerate(header)}
        for r in rows:
            if not r or r[ix["Plant ID"]] in (None, ""):
                continue
            read += 1
            tech = EIA_TECH.get(str(r[ix["Technology"]] or "").strip())
            if not tech:
                continue
            row_status, detail = status, ""
            if status == "planned":
                detail = str(r[ix["Status"]] or "").strip()
                code = re.match(r"\((\w+)\)", detail)
                if code and code.group(1) in UNDER_CONSTRUCTION:
                    row_status = "under_construction"
            elif status == "operating" and "Status" in ix:
                detail = str(r[ix["Status"]] or "").strip()
            pid = _int_text(r[ix["Plant ID"]])
            key = (pid, row_status)
            p = plants.get(key)
            if p is None:
                lat, lon = _num(r[ix["Latitude"]]), _num(r[ix["Longitude"]])
                p = plants[key] = {
                    "facility_id": f"eia:{pid}:{row_status}", "source": "eia860m",
                    "plant_id": pid, "eia_plant_id": pid,
                    "name": str(r[ix["Plant Name"]] or "").strip(),
                    "operator": str(r[ix["Entity Name"]] or "").strip(),
                    "state": str(r[ix["Plant State"]] or "").strip().upper(),
                    "county": str(r[ix["County"]] or "").strip(),
                    "_tech": set(), "_cap": 0.0, "_units": 0, "_dates": [], "_detail": Counter(),
                    "latitude": lat, "longitude": lon,
                }
            p["_tech"].add(tech)
            p["_cap"] += _num(r[ix["Nameplate Capacity (MW)"]]) or 0.0
            p["_units"] += 1
            if detail:
                p["_detail"][detail] += 1
            ycol, mcol = EIA_DATE_COLS.get(status, (None, None))
            if ycol:
                y, mo = _int_text(r[ix[ycol]]), _int_text(r[ix[mcol]])
                if y:
                    p["_dates"].append((y, mo))
    out = []
    for p in plants.values():
        status = p["facility_id"].rsplit(":", 1)[1]
        dates = sorted(p.pop("_dates"), key=lambda d: (int(d[0]), int(d[1] or 0)))
        # In service or retired: the first unit's date. Planned: the last
        # unit's expected date, when the whole plant would be running.
        y, mo = ("", "")
        if dates:
            y, mo = dates[-1] if status in ("planned", "under_construction") else dates[0]
        detail = p.pop("_detail")
        p.update({"technology": ";".join(sorted(p.pop("_tech"))), "capacity_mw": _mw(p.pop("_cap")),
                  "units": p.pop("_units"), "year": y, "month": mo, "status": status,
                  "status_detail": "; ".join(d for d, _ in detail.most_common()),
                  "release": release, "county_fips": ""})
        out.append(p)
    return release, out, read


# ── USWTDB and USPVDB ───────────────────────────────────────────────────────

def _zip_csv(data: bytes, prefix: str) -> tuple[str, list[dict]]:
    """(file name, rows) of the one data CSV in a USGS zip."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise SchemaError(f"{prefix}: not a zip file") from exc
    names = [n for n in zf.namelist() if n.lower().endswith(".csv") and Path(n).name.lower().startswith(prefix)]
    if len(names) != 1:
        raise SchemaError(f"{prefix}: expected one {prefix}*.csv in the zip, found {names}")
    text = zf.read(names[0]).decode("utf-8-sig", errors="replace")
    return Path(names[0]).name, list(csv.DictReader(io.StringIO(text)))


def _release(file_name: str, label: str) -> str:
    m = re.search(r"_[vV](\d+)_(\d+)_(\d{4})(\d{2})(\d{2})", file_name)
    return f"{label} v{m.group(1)}.{m.group(2)} ({m.group(3)}-{m.group(4)}-{m.group(5)})" if m else label


def _require(rows: list[dict], cols: list[str], label: str) -> None:
    if not rows:
        raise SchemaError(f"{label}: no rows")
    lacks = [c for c in cols if c not in rows[0]]
    if lacks:
        raise SchemaError(f"{label}: missing column(s) {lacks}")


def read_uswtdb(data: bytes) -> tuple[str, list[dict], int]:
    """One row per project: turbines grouped by state, project name and EIA id."""
    file_name, rows = _zip_csv(data, "uswtdb")
    _require(rows, USWTDB_REQUIRED, "USWTDB")
    release = _release(file_name, "USWTDB")
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for r in rows:
        groups[(r["t_state"].strip().upper(), r["p_name"].strip(), _int_text(r["eia_id"]))].append(r)
    out = []
    for (state, name, eia), turbines in groups.items():
        if not name or name.lower() in ("unknown", "n/a"):
            continue
        lats = [x for x in (_num(t["ylat"]) for t in turbines) if x is not None]
        lons = [x for x in (_num(t["xlong"]) for t in turbines) if x is not None]
        fips = Counter(t["t_fips"].strip().zfill(5) for t in turbines if t["t_fips"].strip())
        county = Counter(t["t_county"].strip() for t in turbines if t["t_county"].strip())
        years = sorted(y for y in (_int_text(t["p_year"]) for t in turbines) if y and y != "-9999")
        cap = _num(turbines[0]["p_cap"])
        key = eia or hashlib.sha1(f"{state}|{name}".encode()).hexdigest()[:10]
        out.append({
            "facility_id": f"uswtdb:{state}:{key}", "source": "uswtdb", "release": release,
            "plant_id": key, "eia_plant_id": eia, "name": name, "operator": "",
            "state": state, "county": county.most_common(1)[0][0] if county else "",
            "county_fips": fips.most_common(1)[0][0] if fips else "",
            "technology": "wind", "status": "operating", "status_detail": "built (turbines located from imagery)",
            "capacity_mw": _mw(cap) if cap is not None and cap > 0 else "",
            "year": years[0] if years else "", "month": "", "units": len(turbines),
            "latitude": round(sum(lats) / len(lats), 5) if lats else None,
            "longitude": round(sum(lons) / len(lons), 5) if lons else None,
        })
    return release, out, len(rows)


def read_uspvdb(data: bytes) -> tuple[str, list[dict], int]:
    """One row per facility."""
    file_name, rows = _zip_csv(data, "uspvdb")
    _require(rows, USPVDB_REQUIRED, "USPVDB")
    release = _release(file_name, "USPVDB")
    out = []
    for r in rows:
        tech = "battery_storage;solar" if (r.get("p_battery") or "").strip().lower() in ("batteries", "battery", "yes") else "solar"
        cap = _num(r["p_cap_ac"])
        year = _int_text(r["p_year"])
        out.append({
            "facility_id": f"uspvdb:{r['case_id'].strip()}", "source": "uspvdb", "release": release,
            "plant_id": r["case_id"].strip(), "eia_plant_id": _int_text(r["eia_id"]),
            "name": r["p_name"].strip(), "operator": "", "state": r["p_state"].strip().upper(),
            "county": r["p_county"].strip(), "county_fips": "", "technology": tech,
            "status": "operating", "status_detail": "built (digitized from imagery)",
            "capacity_mw": _mw(cap) if cap is not None and cap > 0 else "",
            "year": year if year != "-9999" else "", "month": "", "units": "",
            "latitude": _num(r["ylat"]), "longitude": _num(r["xlong"]),
        })
    return release, out, len(rows)


# ── Assembly ────────────────────────────────────────────────────────────────

def finish(rows: list[dict]) -> list[dict]:
    """Keep the states the pipeline knows, place each plant by its point, sort."""
    out = []
    for r in rows:
        if r["state"] not in STATE_CODES:
            continue
        lat, lon = r.get("latitude"), r.get("longitude")
        r["county_fips"] = _place(lat, lon) or r.get("county_fips") or ""
        r["latitude"] = round(lat, 5) if isinstance(lat, float) else (lat or "")
        r["longitude"] = round(lon, 5) if isinstance(lon, float) else (lon or "")
        out.append(r)
    return sorted(out, key=lambda r: (r["state"], r["source"], r["name"].lower(), r["facility_id"]))


READERS = {"eia860m": read_eia, "uswtdb": read_uswtdb, "uspvdb": read_uspvdb}


def download(session: requests.Session) -> dict[str, tuple[str, bytes]]:
    """source -> (url, bytes). EIA: the newest monthly workbook actually out."""
    out = {}
    resp = session.get(EIA_INDEX, timeout=60)
    resp.raise_for_status()
    for url in eia_candidates(resp.text):
        r = session.get(url, timeout=300)
        if r.ok and r.content.startswith(b"PK"):
            out["eia860m"] = (url, r.content)
            break
    else:
        raise SchemaError("no EIA-860M workbook on the index page could be downloaded")
    for source, url in (("uswtdb", USWTDB_URL), ("uspvdb", USPVDB_URL)):
        r = session.get(url, timeout=300)
        r.raise_for_status()
        out[source] = (url, r.content)
    return out


def build(files: dict[str, tuple[str, bytes]]) -> tuple[list[dict], list[dict]]:
    """(facility rows, manifest rows). Raises SchemaError on any changed layout."""
    rows, manifest = [], []
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for source, (url, data) in files.items():
        release, kept, read = READERS[source](data)
        kept = finish(kept)
        rows += kept
        manifest.append({"source": source, "release": release, "url": url,
                         "file_name": url.rsplit("/", 1)[-1], "sha256": sha256(data), "bytes": len(data),
                         "fetched_at": now, "rows_read": read, "rows_kept": len(kept)})
    return rows, manifest


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--eia", type=Path, help="a local EIA-860M workbook instead of downloading")
    ap.add_argument("--uswtdb", type=Path, help="a local uswtdbCSV.zip")
    ap.add_argument("--uspvdb", type=Path, help="a local uspvdbCSV.zip")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    if args.dry_run:
        print(f"[dry-run] EIA-860M: newest workbook listed on {EIA_INDEX}")
        print(f"[dry-run] USWTDB: {USWTDB_URL}\n[dry-run] USPVDB: {USPVDB_URL}")
        print(f"[dry-run] would write {FACILITIES_PATH.relative_to(ROOT)} and {MANIFEST_PATH.relative_to(ROOT)}")
        return 0
    local = {"eia860m": args.eia, "uswtdb": args.uswtdb, "uspvdb": args.uspvdb}
    files: dict[str, tuple[str, bytes]] = {}
    if any(local.values()):
        if not all(local.values()):
            ap.error("give all three local files, or none")
        urls = {"eia860m": EIA_INDEX + "xls/" + args.eia.name, "uswtdb": USWTDB_URL, "uspvdb": USPVDB_URL}
        files = {source: (urls[source], path.read_bytes()) for source, path in local.items()}
    else:
        session = requests.Session()
        session.headers["User-Agent"] = USER_AGENT
        try:
            files = download(session)
        except (requests.RequestException, SchemaError) as exc:
            print(f"Download failed: {exc}", file=sys.stderr)
            return 1
    try:
        rows, manifest = build(files)
    except SchemaError as exc:
        print(f"Refused by the schema guard, nothing written: {exc}", file=sys.stderr)
        return 1
    write_csv(FACILITIES_PATH, rows, FIELDS)
    write_csv(MANIFEST_PATH, manifest, MANIFEST_FIELDS)
    for m in manifest:
        print(f"{m['source']}: {m['release']}, {m['rows_read']} rows read, {m['rows_kept']} kept "
              f"({m['sha256'][:12]})")
    by = Counter((r["source"], r["status"]) for r in rows)
    print("; ".join(f"{s} {st}: {n}" for (s, st), n in sorted(by.items())))
    print(f"Wrote {FACILITIES_PATH.relative_to(ROOT)} ({len(rows)} rows, "
          f"{sum(1 for r in rows if r['county_fips'])} placed in a county)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
