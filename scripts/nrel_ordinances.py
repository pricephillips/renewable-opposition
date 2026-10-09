"""NREL siting ordinance databases (2025): siting_standards rows and the
restrictions they support.

Sources (stored in data/raw/ by scripts/fetch_nrel_ordinances.py):
  wind   "U.S. Wind Siting Regulation and Zoning Ordinances (2025)", OEDI
         submission 8519, DOI 10.25984/3363758, CC BY 4.0, data from
         2025-09-12, last updated 2026-02-03
  solar  "U.S. Solar Siting Regulation and Zoning Ordinances (2025)", OEDI
         submission 8602, DOI 10.25984/3363739, CC BY 4.0, data from
         2025-12-30, last updated 2026-02-03

NREL compiled both with large language models (INFRA-COMPASS) and says the
data must be validated. Here it is a compiled source: every row is
evidence_level compiled_flagged and verification unverified until a reviewer
reads the ordinance (restriction_sources.csv, docs/AGENT_REVIEW.md).

siting_standards (built from the stored workbooks by build_seed_outputs.py,
through siting_standards.py; there is no seed copy, since it would only repeat
the stored spreadsheets)
  One row per jurisdiction, technology and feature, copied from the "All
  Ordinances" sheet: NREL's state, county, subdivision, jurisdiction type and
  FIPS code, feature, value, units, setback adder, minimum and maximum setback,
  summary, section, ordinance year and source URL. Two NREL rows for the same
  jurisdiction and feature (nine wind cases) become one row with both values,
  ";"-joined. Placement is done by build_seed_outputs.py with the existing
  county logic (classify.county_fips_all); NREL's own county code (the first
  five digits of its FIPS code) is the fallback, method source_fips.

Restrictions (rows of data/seed/restrictions_seed.csv, source "NREL U.S. ...")
  The shared rules only. A jurisdiction and technology becomes a restriction
  when it has an outright prohibition or moratorium, or when
  build_sabin_seeds.restriction_severity scores its combined features at 2 or
  more. Features enter the score through the build_sabin_seeds.MECHANISM_TYPE
  vocabulary (RESTRICTING): setbacks from non-participating structures and
  property lines (the setbacks the severity rule is written for), maximum
  height, noise, maximum project size, maximum lot size, coverage, shadow
  flicker, prohibited use districts and prohibitions. Every other feature
  (setbacks from roads, rails, lines, water and participating owners, turbine
  spacing, blade clearance, decommissioning, lighting, signage, use districts
  and so on) stays in siting_standards only.

  Prohibitions are read from NREL's summary: a moratorium when it says
  moratorium or temporary prohibition; an outright ban when it says the
  facilities are prohibited or not permitted throughout the jurisdiction (all
  districts, anywhere, in the town or county) for utility-scale, commercial
  or principal-use systems; otherwise (some districts only, a sub-type such as
  concentrating solar, a cap or a permit condition) a zoning restriction,
  scored by the default rule.

  status is unknown: NREL does not record whether an ordinance is in force,
  and a moratorium may have lapsed. Severity follows restriction_severity
  (an unknown moratorium scores 4, as for Sabin rows).

Dedupe (data/review/nrel_restriction_conflicts.csv)
  An NREL restriction is matched to Sabin, Moratorium Nation and promoted
  review-queue restrictions by state, jurisdiction (common.jurisdiction_key)
  and technology. Same restriction_type and severity: a duplicate, not
  published as a second restriction (its features stay in siting_standards).
  Otherwise the sources disagree: both are kept, the NREL row notes the
  conflict (conflicts_with) and the pair goes on the conflicts worklist.

Usage
  python scripts/nrel_ordinances.py            rebuild the NREL seed rows
  python scripts/nrel_ordinances.py --dry-run  counts only
"""
from __future__ import annotations

import argparse
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_sabin_seeds import (  # noqa: E402
    MECHANISM_TYPE, RESTRICTION_FIELDS, driving_type, restriction_severity,
)
from common import (  # noqa: E402
    REVIEW_DIR, ROOT, SEED_DIR, STATE_ABBREV, jurisdiction_key, jurisdiction_kind, read_csv, write_csv,
)

RAW_MANIFEST = ROOT / "data" / "raw" / "manifest.csv"
RESTRICTIONS_PATH = SEED_DIR / "restrictions_seed.csv"
CONFLICTS_PATH = REVIEW_DIR / "nrel_restriction_conflicts.csv"

DATASETS = {
    "wind": {
        "source_id": "nrel_wind_ordinances_2025",
        "url": "https://data.openei.org/files/8519/Wind%20Ordinances%202025%20(1).xlsx",
        "label": ("NREL U.S. Wind Siting Regulation and Zoning Ordinances (2025), "
                  "DOI 10.25984/3363758, CC BY 4.0"),
        "short": "NREL wind ordinances (2025), CC BY 4.0",
        "doi_url": "https://doi.org/10.25984/3363758",
        "sheets": {
            "All Ordinances": ["State", "County", "Subdivision", "Jurisdiction Type",
                               "County Subdivision FIPS Code", "Feature", "Value", "Units", "Summary",
                               "Setback Adder", "Minimum Setback Distance", "Maximum Setback Distance",
                               "Section", "Ordinance Year", "Source URL"],
            "State-Level": ["State", "County Subdivision FIPS Code", "Feature", "Value", "Units", "Summary",
                            "Section", "Ordinance Year", "Source URL"],
            "County-Level": ["State", "County", "County Subdivision FIPS Code", "Feature", "Value", "Units",
                             "Summary", "Setback Adder", "Minimum Setback Distance", "Maximum Setback Distance",
                             "Section", "Ordinance Year", "Source URL"],
            "Municipal-Level": ["State", "County", "Subdivision", "County Subdivision FIPS Code", "Feature",
                                "Value", "Units", "Summary", "Setback Adder", "Minimum Setback Distance",
                                "Maximum Setback Distance", "Section", "Ordinance Year", "Source URL"],
            "Column Descriptions": ["Column", "Description"],
            "Ordinance Types": ["Feature", "Description"],
        },
    },
    "solar": {
        "source_id": "nrel_solar_ordinances_2025",
        "url": "https://data.openei.org/files/8602/Solar%20Ordinances%202025.xlsx",
        "label": ("NREL U.S. Solar Siting Regulation and Zoning Ordinances (2025), "
                  "DOI 10.25984/3363739, CC BY 4.0"),
        "short": "NREL solar ordinances (2025), CC BY 4.0",
        "doi_url": "https://doi.org/10.25984/3363739",
        "sheets": {
            "All Ordinances": ["State", "County", "Subdivision", "Jurisdiction Type",
                               "County Subdivision FIPS Code", "Feature", "Value", "Units", "Summary",
                               "Setback Adder", "Section", "Ordinance Year", "Source URL"],
            "State-Level": ["State", "County Subdivision FIPS Code", "Feature", "Value", "Units", "Summary",
                            "Section", "Ordinance Year", "Source URL"],
            "County-Level": ["State", "County", "County Subdivision FIPS Code", "Feature", "Value", "Units",
                             "Summary", "Setback Adder", "Section", "Ordinance Year", "Source URL"],
            "Municipal-Level": ["State", "County", "Subdivision", "County Subdivision FIPS Code", "Feature",
                                "Value", "Units", "Summary", "Setback Adder", "Section", "Ordinance Year",
                                "Source URL"],
            "Column Descriptions": ["Column", "Description"],
            "Ordinance Types": ["Feature", "Description"],
        },
    },
}
SOURCE_PREFIX = "NREL U.S. "
DISCLAIMER = "This data was collected with the help of generative AI"
# NREL's feature names, per technology. A feature outside these stops the run.
FEATURES = {
    "wind": {"Accessory Use Districts", "Blade Clearance", "Climbing Prevention", "Color", "Decommissioning",
             "Lighting", "Maximum Height", "Maximum Lot Size", "Maximum Project Size", "Minimum Lot Size",
             "Noise", "Other Wecs", "Primary Use Districts", "Prohibitions", "Property Line (Non-Participating)",
             "Property Line (Participating)", "Railroads", "Repowering", "Roads", "Shadow Flicker", "Signage",
             "Soil", "Special Use Districts", "Structures (Non-Participating)", "Structures (Participating)",
             "Tower Density", "Transmission", "Visual Impact", "Water"},
    "solar": {"Accessory Use Districts", "Coverage", "Decommissioning", "Fencing", "Glare", "Maximum Height",
              "Maximum Lot Size", "Maximum Project Size", "Minimum Lot Size", "Noise", "Other Secs",
              "Panel Spacing", "Primary Use Districts", "Prohibited Use Districts", "Prohibitions",
              "Property Line (Non-Participating)", "Property Line (Participating)", "Public Conservation Lands",
              "Railroads", "Repowering", "Roads", "Screening", "Signage", "Soil", "Special Use Districts",
              "Structures (Non-Participating)", "Structures (Participating)", "Transmission", "Visual Impact",
              "Water"},
}
# Features that enter the restriction score, in the MECHANISM_TYPE vocabulary.
# Prohibitions are classified from the summary (prohibition_mechanism).
RESTRICTING = {
    "Structures (Non-Participating)": "setback",
    "Property Line (Non-Participating)": "setback",
    "Maximum Height": "height limit",
    "Noise": "noise limit",
    "Maximum Project Size": "cap on project size",
    "Maximum Lot Size": "cap on area",
    "Coverage": "cap on area",
    "Shadow Flicker": "shadow_flicker_limit",
    "Prohibited Use Districts": "zoning restriction",
    "Prohibitions": None,
}
COUNTY_TYPES = {"county", "parish"}
# The restriction threshold. restriction_severity never scores a restricting
# mechanism below 2, so at 2 every jurisdiction with a restricting feature is
# a restriction; 3 would keep only the severe ones (and every prohibition).
MIN_SEVERITY = 2

SITING_FIELDS = [
    "standard_id", "state", "jurisdiction", "jurisdiction_type", "nrel_jurisdiction_type", "county_name",
    "subdivision", "technology", "feature", "value", "units", "setback_adder", "min_setback_ft",
    "max_setback_ft", "summary", "section", "ordinance_year", "ordinance_url", "nrel_fips_code",
    "source_county_fips", "restricting", "nrel_id", "source", "source_url",
]
NREL_RESTRICTION_FIELDS = RESTRICTION_FIELDS + ["nrel_id", "ordinance_url", "source_county_fips",
                                                "nrel_jurisdiction_type", "conflicts_with"]
CONFLICT_FIELDS = ["nrel_id", "state", "jurisdiction", "technology", "nrel_restriction_type", "nrel_severity",
                   "nrel_basis", "other_instrument", "other_source", "other_restriction_type",
                   "other_severity", "decision", "ordinance_url"]


class SchemaError(SystemExit):
    """The spreadsheet's structure is not what this code reads."""


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ── Reading and the schema guard ─────────────────────────────────────────────

def _header(row: tuple) -> list[str]:
    cells = [_s(c) for c in row]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def sheet_rows(ws) -> tuple[list[str], list[tuple]]:
    """(header, data rows) of a sheet, skipping NREL's disclaimer line."""
    rows = list(ws.iter_rows(values_only=True))
    start = 1 if rows and _s(rows[0][0]).startswith(DISCLAIMER) else 0
    return (_header(rows[start]) if len(rows) > start else []), rows[start + 1:]


def check_schema(data: bytes, tech: str) -> list[str]:
    """Problems with the workbook's structure: a missing sheet, a changed
    header, an unknown feature. Empty when it is what this code expects."""
    import openpyxl
    expected = DATASETS[tech]["sheets"]
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    except Exception as exc:  # not a workbook at all
        return [f"not a readable .xlsx: {exc}"]
    problems = []
    for name, header in expected.items():
        if name not in wb.sheetnames:
            problems.append(f"sheet {name!r} is missing (sheets: {wb.sheetnames})")
            continue
        got, rows = sheet_rows(wb[name])
        if got != header:
            problems.append(f"sheet {name!r} header changed: expected {header}, got {got}")
        elif name == "All Ordinances":
            feats = {_s(r[header.index("Feature")]) for r in rows if any(c is not None for c in r)}
            unknown = sorted(feats - FEATURES[tech] - {""})
            if unknown:
                problems.append(f"unknown feature(s) {unknown}; map them in FEATURES and RESTRICTING")
    extra = sorted(set(wb.sheetnames) - set(expected))
    if extra:
        problems.append(f"unexpected sheet(s) {extra}")
    return problems


def newest_raw(tech: str) -> Path:
    sid = DATASETS[tech]["source_id"]
    stored = [r for r in read_csv(RAW_MANIFEST) if r.get("source_id") == sid and not r.get("error")
              and r.get("local_path") not in ("", "unchanged")]
    if not stored:
        raise SystemExit(f"No stored copy of {sid}; run scripts/fetch_nrel_ordinances.py")
    return ROOT / stored[-1]["local_path"]


def load(tech: str, path: Path | None = None) -> list[dict]:
    """The All Ordinances rows as dicts, after the schema guard."""
    import openpyxl
    data = (path or newest_raw(tech)).read_bytes()
    problems = check_schema(data, tech)
    if problems:
        raise SchemaError(f"{tech} ordinances: " + "; ".join(problems))
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    header, rows = sheet_rows(wb["All Ordinances"])
    return [dict(zip(header, r)) for r in rows if any(c is not None for c in r[:len(header)])]


# ── Rows ─────────────────────────────────────────────────────────────────────

def fips_parts(code) -> tuple[str, str]:
    """(NREL code as text, the county FIPS it names or ''). County codes are
    5 digits (1003 -> 01003); a municipality's code is the county code
    followed by 5 digits; a state's ends in 000."""
    c = _s(code)
    if not c.isdigit():
        return c, ""
    if len(c) <= 5:
        c5 = c.zfill(5)
        return c5, "" if c5.endswith("000") else c5
    c10 = c.zfill(10)
    return c10, c10[:5]


def state_of(name: str) -> str:
    return STATE_ABBREV.get(_s(name).lower(), "")


def jurisdiction(r: dict) -> tuple[str, str, str]:
    """(name, our jurisdiction_type, NREL's type). Counties keep their name
    with County or Parish; towns, townships, cities and boroughs are named
    '<Subdivision> <Type>' and typed Municipality (a Pennsylvania borough is
    not a county). State-level rows are type State."""
    nrel_type = _s(r.get("Jurisdiction Type"))
    county, sub = _s(r.get("County")), _s(r.get("Subdivision"))
    if not nrel_type and not county and not sub:
        return STATE_NAME.get(state_of(r["State"]), _s(r["State"])), "State", ""
    if nrel_type.lower() in COUNTY_TYPES or (not sub and county):
        word = "Parish" if nrel_type.lower() == "parish" else "County"
        return f"{county} {word}", word, nrel_type
    name = sub
    if nrel_type and not re.search(rf"\b{re.escape(nrel_type)}\b", sub, re.I):
        name = f"{sub} {nrel_type}"
    return name, "Municipality", nrel_type


STATE_NAME = {v: k.title() for k, v in STATE_ABBREV.items()}


def standard_rows(tech: str, raw: list[dict]) -> list[dict]:
    meta = DATASETS[tech]
    out: dict[tuple, dict] = {}
    for r in raw:
        st = state_of(r["State"])
        if not st:
            raise SchemaError(f"{tech}: unknown state {r['State']!r}")
        name, jtype, nrel_type = jurisdiction(r)
        code, county_fips = fips_parts(r.get("County Subdivision FIPS Code"))
        feature = _s(r["Feature"])
        row = {
            "state": st, "jurisdiction": name, "jurisdiction_type": jtype, "nrel_jurisdiction_type": nrel_type,
            "county_name": _s(r.get("County")), "subdivision": _s(r.get("Subdivision")), "technology": tech,
            "feature": feature, "value": _s(r.get("Value")), "units": _s(r.get("Units")),
            "setback_adder": _s(r.get("Setback Adder")), "min_setback_ft": _s(r.get("Minimum Setback Distance")),
            "max_setback_ft": _s(r.get("Maximum Setback Distance")), "summary": _s(r.get("Summary")),
            "section": _s(r.get("Section")), "ordinance_year": _s(r.get("Ordinance Year")),
            "ordinance_url": " ".join(_s(r.get("Source URL")).split()),
            "nrel_fips_code": code, "source_county_fips": county_fips,
            "restricting": "yes" if feature in RESTRICTING else "",
            "nrel_id": f"{tech}:{code}", "source": meta["short"], "source_url": meta["doi_url"],
        }
        key = (tech, code, st, name, feature)
        if key in out:
            prev = out[key]
            for f in ("value", "units", "summary", "section", "setback_adder", "min_setback_ft",
                      "max_setback_ft", "ordinance_url"):
                if row[f] and row[f] != prev[f]:
                    prev[f] = "; ".join(x for x in (prev[f], row[f]) if x)
            continue
        row["standard_id"] = f"nrel:{tech}:{code}:{re.sub(r'[^a-z0-9]+', '_', feature.lower()).strip('_')}"
        out[key] = row
    return sorted(out.values(), key=lambda x: (x["state"], x["jurisdiction"], x["technology"], x["feature"]))


# ── Restrictions ─────────────────────────────────────────────────────────────

_MORATORIUM = re.compile(r"moratori|temporar(?:y|ily) prohibit|temporarily (?:suspend|halt)", re.I)
_OUTRIGHT = re.compile(
    r"prohibited (?:uses? )?in all|not permitted in any|(?:is|are) not (?:a )?permitted (?:use )?in (?:any|all)"
    r"|prohibited (?:anywhere|throughout|within the (?:entire|county|town|township|city|village|borough))"
    r"|not (?:be )?(?:allowed|permitted) (?:anywhere|in the (?:town|township|county|city|village|borough))"
    r"|not permitted within (?:the )?(?:county|town|township|city|village|borough|jurisdiction)"
    r"|in all of the unincorporated|prohibit(?:s|ed)? .{0,80}in all zoning districts"
    r"|not considered to be an appropriate land use in the county", re.I)
_SUBTYPE = re.compile(r"concentrat|\bcsp\b|roof|pole[- ]mounted|accessory|residential (?:solar|wind) (?:energy )?"
                      r"system|small(?:[- ]scale)? (?:wind|solar)|private wecs", re.I)
_SCALE = re.compile(r"utility|commercial|large|principal|farm|industrial|grid|wecs|wind energy (?:facilit|system|"
                    r"conversion)|solar energy (?:facilit|system)|megawatt|\bmw\b|wind turbines", re.I)
_UNLESS = re.compile(r"\bunless\b|\bexcept\b|special (?:use|exception)|conditional use", re.I)


def prohibition_mechanism(summary: str) -> str:
    """'moratorium', 'ban/prohibition' or 'zoning restriction' (module docstring)."""
    if _MORATORIUM.search(summary):
        return "moratorium"
    if _OUTRIGHT.search(summary) and _SCALE.search(summary) and not _SUBTYPE.search(summary) \
            and not _UNLESS.search(summary.split(".")[0]):
        return "ban/prohibition"
    return "zoning restriction"


_UNIT_WORDS = {
    "feet": "{v} feet", "miles": "{v} miles", "tip-height-multiplier": "{v} times tip height",
    "hub-height-multiplier": "{v} times hub height", "structure-height-multiplier": "{v} times structure height",
    "rotor-diameter-multiplier": "{v} rotor diameters",
}


def feature_text(s: dict) -> str:
    """The feature as words the shared severity rule reads (setback_feet,
    height_multiplier, min_noise_dba). Only A-weighted noise is written as
    dBA, and only setbacks as "<n> feet": other noise units are spelled out
    and a height limit's unit is put in brackets, so neither is read as a
    noise level or a setback distance."""
    v, u, f = s["value"], s["units"], s["feature"]
    if f == "Noise":
        if u == "dBA" and v:
            return f"noise limit {v} dBA"
        words = {"dB": "decibels", "dBC": "decibels, C-weighted", "dBH": "decibels, H-weighted",
                 "dB above background": "decibels above background"}.get(u, u)
        return f"noise limit {v} ({words})" if v else "noise limit"
    if RESTRICTING.get(f) == "setback":
        parts = [_UNIT_WORDS.get(u, "{v} " + u).format(v=v)] if v else []
        if s["min_setback_ft"]:
            parts.append(f"at least {s['min_setback_ft']} feet")
        return f"setback from {f.lower()} " + ", ".join(parts)
    if f == "Maximum Height":
        # "(in feet)", never "430 feet": setback_feet reads any "<n> feet".
        return f"height limit {v} (in {u})" if v and u else "height limit"
    if f in ("Maximum Project Size", "Maximum Lot Size", "Coverage"):
        return f"{f.lower()} {v} {u}".strip()
    if f == "Shadow Flicker":
        return f"shadow flicker limit {v} {u.replace('hr/year', 'hours per year')}".strip()
    return f"{f.lower()}"


def restriction_rows(standards: list[dict]) -> list[dict]:
    """One candidate restriction per jurisdiction and technology whose
    restricting features score 2 or more (module docstring)."""
    groups: dict[str, list[dict]] = {}
    for s in standards:
        if s["restricting"] and s["jurisdiction_type"] != "State":
            groups.setdefault(s["nrel_id"], []).append(s)
    out = []
    for nid, feats in sorted(groups.items()):
        mechs = []
        for s in feats:
            m = prohibition_mechanism(s["summary"]) if s["feature"] == "Prohibitions" else RESTRICTING[s["feature"]]
            if m not in mechs:
                mechs.append(m)
        types = {MECHANISM_TYPE[m] for m in mechs}
        text = "; ".join(feature_text(s) for s in feats if s["feature"] != "Prohibitions")
        first = feats[0]
        tech = first["technology"]
        score, basis = restriction_severity(types, tech, "unknown", text)
        if score < MIN_SEVERITY and not types & {"ban", "moratorium"}:
            continue
        values = "; ".join(feature_text(s) for s in feats if s["feature"] != "Prohibitions")
        prohibitions = [s for s in feats if s["feature"] == "Prohibitions"]
        desc = f"{first['jurisdiction']} {tech} siting ordinance, as compiled by NREL: " + (
            values + ("; " if values and prohibitions else "") +
            ("; ".join(f"prohibition ({prohibition_mechanism(p['summary'])})" for p in prohibitions)))
        urls = sorted({s["ordinance_url"] for s in feats if s["ordinance_url"]})
        years = sorted({s["ordinance_year"] for s in feats if s["ordinance_year"]})
        out.append({
            "state": first["state"], "technology": tech, "restriction_type": driving_type(types, basis),
            "severity_score": score, "description": desc[:400], "status": "unknown",
            "jurisdiction": first["jurisdiction"], "jurisdiction_type": first["jurisdiction_type"],
            "date_enacted_iso": "", "date_text": ", ".join(years), "mechanisms": ", ".join(mechs),
            "severity_basis": basis,
            # NREL's summary of each feature is in siting_standards (standard_id
            # nrel:<technology>:<code>:<feature>); not repeated here.
            "long_description": "",
            "moratorium_id": "", "source_record_id": "", "source": DATASETS[tech]["label"],
            "source_url": DATASETS[tech]["doi_url"],
            "notes": ("NREL compiled this with large language models and asks that it be validated; "
                      "not verified against the ordinance. NREL does not record whether it is in force."),
            "nrel_id": nid, "ordinance_url": " ".join(urls), "source_county_fips": first["source_county_fips"],
            "nrel_jurisdiction_type": first["nrel_jurisdiction_type"], "conflicts_with": "",
        })
    return out


def other_index(existing: list[dict]) -> dict[tuple[str, str], list[dict]]:
    """(jurisdiction_key, technology) -> restriction seed rows from Sabin,
    Moratorium Nation and the review queue (every non-NREL row)."""
    idx: dict[tuple[str, str], list[dict]] = {}
    for r in existing:
        if (r.get("source") or "").startswith(SOURCE_PREFIX) or not r.get("state"):
            continue
        kind = jurisdiction_kind(r.get("jurisdiction_type", ""))
        if kind == "state":
            continue
        idx.setdefault((jurisdiction_key(r["state"], kind, r.get("jurisdiction", "")), r["technology"]), []).append(r)
    return idx


def _instrument(r: dict) -> str:
    import classify
    return classify.instrument_id("restrictions", r)


def dedupe(candidates: list[dict], existing: list[dict]) -> tuple[list[dict], list[dict]]:
    """(rows to publish, conflict worklist rows), per the module docstring."""
    idx = other_index(existing)
    keep, conflicts = [], []
    for c in candidates:
        kind = jurisdiction_kind(c["jurisdiction_type"])
        others = idx.get((jurisdiction_key(c["state"], kind, c["jurisdiction"]), c["technology"]), [])
        if not others:
            keep.append(c)
            continue
        same = [o for o in others if o.get("restriction_type") == c["restriction_type"]
                and str(o.get("severity_score")) == str(c["severity_score"])]
        decision = "duplicate: not published again" if same else "conflict: both kept, review"
        for o in (same[:1] if same else others):
            conflicts.append({
                "nrel_id": c["nrel_id"], "state": c["state"], "jurisdiction": c["jurisdiction"],
                "technology": c["technology"], "nrel_restriction_type": c["restriction_type"],
                "nrel_severity": c["severity_score"], "nrel_basis": c["severity_basis"],
                "other_instrument": _instrument(o), "other_source": (o.get("source") or "")[:60],
                "other_restriction_type": o.get("restriction_type", ""),
                "other_severity": o.get("severity_score", ""), "decision": decision,
                "ordinance_url": c["ordinance_url"]})
        if not same:
            c["conflicts_with"] = "; ".join(sorted({_instrument(o) for o in others}))
            keep.append(c)
    return keep, conflicts


def all_standards() -> list[dict]:
    """Every siting_standards row, from the newest stored workbooks, with
    nrel_restriction set for the features of a published NREL restriction."""
    standards = []
    for tech in DATASETS:
        standards += standard_rows(tech, load(tech))
    published = {r["nrel_id"] for r in read_csv(RESTRICTIONS_PATH) if r.get("nrel_id")}
    for s in standards:
        s["nrel_restriction"] = "yes" if s["nrel_id"] in published and s["restricting"] else ""
    return standards


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    standards = []
    for tech in DATASETS:
        standards += standard_rows(tech, load(tech))
    existing = read_csv(RESTRICTIONS_PATH)
    kept = [r for r in existing if not (r.get("source") or "").startswith(SOURCE_PREFIX)]
    candidates = restriction_rows(standards)
    rows, conflicts = dedupe(candidates, kept)
    published = {r["nrel_id"] for r in rows}
    for s in standards:
        s["nrel_restriction"] = "yes" if s["nrel_id"] in published and s["restricting"] else ""
    by_sev: dict[int, int] = {}
    for r in rows:
        by_sev[r["severity_score"]] = by_sev.get(r["severity_score"], 0) + 1
    dup = sum(1 for c in conflicts if c["decision"].startswith("duplicate"))
    print(f"siting_standards: {len(standards)} rows ({sum(1 for s in standards if s['technology'] == 'wind')} wind, "
          f"{sum(1 for s in standards if s['technology'] == 'solar')} solar) in "
          f"{len({s['nrel_id'] for s in standards})} jurisdiction-technology pairs")
    print(f"NREL restrictions: {len(candidates)} candidates; {len(rows)} published "
          f"(severity {', '.join(f'{k}={v}' for k, v in sorted(by_sev.items()))}); {dup} duplicates of Sabin, "
          f"Moratorium Nation or queue rows; {len({c['nrel_id'] for c in conflicts}) - dup} kept with a conflict")
    if args.dry_run:
        return 0
    header = list(existing[0].keys()) if existing else []
    write_csv(RESTRICTIONS_PATH, kept + rows, header + [f for f in NREL_RESTRICTION_FIELDS if f not in header])
    write_csv(CONFLICTS_PATH, conflicts, CONFLICT_FIELDS)
    print(f"Wrote {RESTRICTIONS_PATH.relative_to(ROOT)}, {CONFLICTS_PATH.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
