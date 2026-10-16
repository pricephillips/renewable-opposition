"""Render a case-study markdown file into a .docx built on the Coweta example.

The Coweta document is the template: its body is emptied and refilled, so the
page setup, styles, footers (with the Hawthorn mark), list numbering and the
anchored first-page map are its own. The markdown dialect is small:

  # / ##                    Heading 1 / Heading 2
  - item / "  - item"       bullet, level 0 / 1 (Key Takeaways use Coweta's numId 11)
  1. item                   numbered list; each block restarts at 1
  | a | b |                 table; the first table is the ratings table (Plain Table 5)
  ::widths 2700,2500,4880:: column widths in twips for the next table
  ::map path.png 3.0::      first-page map, anchored on the right of the next paragraph
  ::small::                 following paragraphs set in 8 pt (sources, credits)
  [^key] / [^key]: text     footnotes
  **b** *i* ***bi*** [t](url) ^sup^

Usage: python render_docx.py IN.md OUT.docx TEMPLATE.docx TITLE
"""
import copy
import re
import sys

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Emu
from lxml import etree
from PIL import Image

SRC, OUT, TEMPLATE, TITLE = sys.argv[1:5]
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

d = docx.Document(TEMPLATE)
body = d.element.body
# Keep the anchored map drawing from the template before emptying the body.
map_drawing = copy.deepcopy(d.paragraphs[2]._p.find(".//" + qn("w:drawing")))
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

# Drop relationships the old body used: images, hyperlinks, comments.
for rid, rel in list(d.part.rels.items()):
    kind = rel.reltype.split("/")[-1]
    if rel.is_external or kind in ("image", "comments", "commentsExtended", "commentsIds",
                                   "commentsExtensible", "people"):
        del d.part.rels[rid]

numbering = d.part.numbering_part.element
BULLET_TAKEAWAY, BULLET, NUMBERED = "11", "9", "5"


def new_num(base_num_id):
    """A fresh w:num on the same abstract list as base_num_id, restarting at 1."""
    base = numbering.find(f"{{{W}}}num[@{{{W}}}numId='{base_num_id}']")
    abstract = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    nid = str(max(ids) + 1)
    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{abstract}"/>'
                    f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>')
    numbering.append(num)
    return nid


# Footnotes part (python-docx keeps it as a raw part)
fn_part = [r.target_part for r in d.part.rels.values() if r.reltype.endswith("/footnotes")][0]
fn_root = etree.fromstring(fn_part.blob)
for f in fn_root.findall(qn("w:footnote")):
    if f.get(qn("w:type")) not in ("separator", "continuationSeparator"):
        fn_root.remove(f)
footnote_ids = {}
footnote_text = {}

text = open(SRC, encoding="utf-8").read()
# footnote definitions
for m in re.finditer(r"^\[\^([^\]]+)\]:\s*(.+)$", text, re.M):
    footnote_text[m.group(1)] = m.group(2).strip()
text = re.sub(r"^\[\^([^\]]+)\]:.*$\n?", "", text, flags=re.M)

INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*\s][^*]*?\*|\[\^[^\]]+\]|\[[^\]]+\]\([^)]+\)|\^[^^\s]+\^|https?://[^\s)<>]+[^\s)<>.,;])")


def add_run(p, s, bold=False, italic=False, size=None, sup=False, style=None):
    r = p.add_run(s)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    if size:
        r.font.size = Pt(size)
    if sup:
        r.font.superscript = True
    if style:
        r.style = style
    return r


def add_link(p, label, url, size=None, bold=False, italic=False):
    rid = d.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    h.set(qn("w:history"), "1")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rs = OxmlElement("w:rStyle")
    rs.set(qn("w:val"), "Hyperlink")
    rPr.append(rs)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if italic:
        rPr.append(OxmlElement("w:i"))
    if size:
        for tag in ("w:sz", "w:szCs"):
            e = OxmlElement(tag)
            e.set(qn("w:val"), str(int(size * 2)))
            rPr.append(e)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = label
    r.append(t)
    h.append(r)
    p._p.append(h)


def add_footnote(p, key, size=None):
    if key not in footnote_ids:
        fid = len(footnote_ids) + 1
        footnote_ids[key] = fid
        body_text = footnote_text[key]
        fn = parse_xml(
            f'<w:footnote {nsdecls("w")} w:id="{fid}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
            f'<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/><w:rFonts w:ascii="Georgia" w:hAnsi="Georgia"/>'
            f'<w:sz w:val="16"/><w:szCs w:val="16"/></w:rPr><w:footnoteRef/></w:r>'
            f'<w:r><w:rPr><w:rFonts w:ascii="Georgia" w:hAnsi="Georgia"/><w:sz w:val="16"/><w:szCs w:val="16"/></w:rPr>'
            f'<w:t xml:space="preserve"> </w:t></w:r></w:p></w:footnote>')
        fn.find(".//" + qn("w:t")).text = " " + re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", body_text).replace("**", "").replace("*", "")
        fn_root.append(etree.fromstring(etree.tostring(fn)))
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rs = OxmlElement("w:rStyle")
    rs.set(qn("w:val"), "FootnoteReference")
    rPr.append(rs)
    r.append(rPr)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(footnote_ids[key]))
    r.append(ref)
    p._p.append(r)


def inline(p, s, size=None, bold=False, italic=False):
    pos = 0
    for m in INLINE.finditer(s):
        if m.start() > pos:
            add_run(p, s[pos:m.start()], bold=bold, italic=italic, size=size)
        tok = m.group(0)
        if tok.startswith("***"):
            inline(p, tok[3:-3], size=size, bold=True, italic=True)
        elif tok.startswith("**"):
            inline(p, tok[2:-2], size=size, bold=True, italic=italic)
        elif tok.startswith("[^"):
            add_footnote(p, tok[2:-1], size=size)
        elif tok.startswith("["):
            lab, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok).groups()
            add_link(p, lab, url, size=size, bold=bold, italic=italic)
        elif tok.startswith("http"):
            add_link(p, tok, tok, size=size, bold=bold, italic=italic)
        elif tok.startswith("^"):
            add_run(p, tok[1:-1], bold=bold, italic=italic, size=size, sup=True)
        else:
            inline(p, tok[1:-1], size=size, bold=bold, italic=True)
        pos = m.end()
    if pos < len(s):
        add_run(p, s[pos:], bold=bold, italic=italic, size=size)


def set_num(p, num_id, level=0):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), str(level))
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), num_id)
    numPr.append(il)
    numPr.append(ni)
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is not None:
        pStyle.addnext(numPr)
    else:
        pPr.insert(0, numPr)


def cell_props(cell, width, header, plain):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(width))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)
    if not plain:
        b = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "BFBFBF")):
                e.set(qn(k), v)
            b.append(e)
        tcPr.append(b)
        if header:
            shd = OxmlElement("w:shd")
            for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "E7E6E6")):
                shd.set(qn(k), v)
            tcPr.append(shd)
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", 30), ("left", 80), ("bottom", 30), ("right", 80)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(w))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


table_count = 0


def add_table(rows, widths):
    global table_count
    plain = table_count == 0
    table_count += 1
    ncol = len(rows[0])
    if not widths:
        lens = [max(len(re.sub(r"[*\[\]]", "", r[c])) for r in rows) for c in range(ncol)]
        lens = [min(max(l, 6), 60) for l in lens]
        widths = [int(10080 * l / sum(lens)) for l in lens]
    t = d.add_table(rows=len(rows), cols=ncol)
    tblPr = t._tbl.tblPr
    for child in list(tblPr):
        tblPr.remove(child)
    for tag, attrs in (("w:tblStyle", {"w:val": "PlainTable5" if plain else "TableNormal"}),
                       ("w:tblW", {"w:w": str(sum(widths)), "w:type": "dxa"}),
                       ("w:tblLayout", {"w:type": "fixed"}),
                       ("w:tblLook", {"w:val": "04A0", "w:firstRow": "1", "w:lastRow": "0",
                                      "w:firstColumn": "1", "w:lastColumn": "0", "w:noHBand": "0",
                                      "w:noVBand": "1"})):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        tblPr.append(e)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for ri, row in enumerate(rows):
        tr = t.rows[ri]
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell_props(cell, widths[ci], ri == 0, plain)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if ri == 0 or len(rows) <= 14 and ri < len(rows) - 1:
                p.paragraph_format.keep_with_next = True
            size = 10 if plain else 9
            inline(p, val.strip(), size=size, bold=(ri == 0))
    return t


def map_paragraph(p, path, width_in):
    img = Image.open(path)
    w, h = img.size
    cx = int(width_in * 914400)
    cy = int(cx * h / w)
    rid, _ = d.part.get_or_add_image(path)
    drawing = copy.deepcopy(map_drawing)
    for e in drawing.iter():
        if e.tag.endswith("}extent") or e.tag.endswith("}ext") and e.get("cx"):
            e.set("cx", str(cx))
            e.set("cy", str(cy))
        if e.tag.endswith("}blip"):
            e.set(qn("r:embed"), rid)
        if e.tag.endswith("}docPr") or e.tag.endswith("}cNvPr"):
            e.set("descr", TITLE + " locator map")
        if e.tag.endswith("}positionH"):
            off = e.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}posOffset")
            off.text = str(int(7.0 * 914400 - cx))
    r = p.add_run()
    r._r.append(drawing)


lines = text.splitlines()
i = 0
small = False
pending_widths = None
pending_map = None
num_for_block = None
in_takeaways = False
while i < len(lines):
    line = lines[i]
    s = line.strip()
    if not s:
        num_for_block = None
        i += 1
        continue
    m = re.match(r"::(\w+)\s*(.*?)::$", s)
    if m:
        cmd, arg = m.groups()
        if cmd == "widths":
            pending_widths = [int(x) for x in arg.split(",")]
        elif cmd == "map":
            path, wi = arg.split()
            pending_map = (path, float(wi))
        elif cmd == "small":
            small = True
        elif cmd == "pagebreak":
            d.add_paragraph().add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
        i += 1
        continue
    if s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c for c in re.split(r"(?<!\\)\|", lines[i].strip())[1:-1]]
            if not all(re.fullmatch(r"\s*:?-+:?\s*", c) for c in cells):
                rows.append([c.replace("\\|", "|") for c in cells])
            i += 1
        add_table(rows, pending_widths)
        pending_widths = None
        nxt = d.add_paragraph()
        nxt.paragraph_format.space_after = Pt(2)
        continue
    if s.startswith("# "):
        p = d.add_paragraph(style="Heading 1")
        inline(p, s[2:])
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        small = False
        in_takeaways = False
        i += 1
        continue
    if s.startswith("## "):
        p = d.add_paragraph(style="Heading 2")
        inline(p, s[3:])
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        in_takeaways = s[3:].strip().lower() == "key takeaways"
        small = s[3:].strip().lower() in ("sources", "image credits")
        i += 1
        continue
    mb = re.match(r"^(\s*)- (.*)$", line)
    mn = re.match(r"^(\s*)\d+\. (.*)$", line)
    if mb or mn:
        level = 1 if len((mb or mn).group(1)) >= 2 else 0
        p = d.add_paragraph(style="List Paragraph")
        if mb:
            set_num(p, BULLET_TAKEAWAY if in_takeaways else BULLET, level)
        else:
            if num_for_block is None:
                num_for_block = new_num(NUMBERED)
            set_num(p, num_for_block, 0)
        p.paragraph_format.space_after = Pt(4)
        inline(p, (mb or mn).group(2), size=8 if small else None)
        i += 1
        continue
    mc = re.match(r"^Table (\d+): (.*)$", s)
    if mc:
        p = d.add_paragraph()
        add_run(p, f"Table {mc.group(1)} ", bold=True, italic=True)
        inline(p, mc.group(2), italic=True)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        i += 1
        continue
    # paragraph (single line)
    p = d.add_paragraph()
    if pending_map:
        map_paragraph(p, *pending_map)
        pending_map = None
    inline(p, s, size=8 if small else None)
    if re.match(r"^\*\*\*Table \d", s):
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
    else:
        p.paragraph_format.space_after = Pt(2 if small else 10)
    i += 1

missing = set(footnote_text) - set(footnote_ids)
if missing:
    print("WARNING unreferenced footnotes:", missing)
fn_part._blob = etree.tostring(fn_root, xml_declaration=True, encoding="UTF-8", standalone=True)
cp = d.core_properties
cp.title = TITLE
cp.author = "Hawthorn Intelligence"
cp.comments = ""
d.save(OUT)
print("wrote", OUT, "tables", table_count, "footnotes", len(footnote_ids))
